"""Filtered Excel/PDF reports, written atomically without exposing user accounts."""
import os
import tempfile
import textwrap
from datetime import datetime
from html import escape
from pathlib import Path

from backend.dashboard import MODULES, export_records


class ExportError(ValueError):
    pass


def report_details(module_key, search, status, count):
    return (f'PersonalManager · {MODULES[module_key].title}',
            f'Export du {datetime.now():%d/%m/%Y %H:%M} · {count} enregistrement(s)',
            f'Recherche : {search.strip() or "Toutes"} · Statut : {status or "Tous"}')


def write_excel(path, module_key, rows, details):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MODULES[module_key].title
    for number, text in enumerate(details, 1):
        sheet.cell(number, 1, text).data_type = 's'
        sheet.merge_cells(start_row=number, start_column=1, end_row=number,
                          end_column=len(MODULES[module_key].columns))
        sheet.row_dimensions[number].height = 30 if number == 1 else 24
        sheet.cell(number, 1).alignment = Alignment(wrap_text=True, vertical='center')
    sheet['A1'].font = Font(name='Calibri', size=18, bold=True, color='142033')
    module = MODULES[module_key]
    sheet.append(list(module.headings))  # Row 4, below report metadata.
    for cell in sheet[4]:
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='142033')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[4].height = 30
    for row_number, row in enumerate(rows, 5):
        lines = 1
        for column_number, (key, value) in enumerate(zip(module.columns, row), 1):
            cell = sheet.cell(row_number, column_number)
            # IDs/phones and integers beyond Excel's 15 digits remain exact text.
            if key == 'amount' and isinstance(value, int) and len(str(abs(value))) <= 15:
                cell.value = value
                cell.number_format = '#,##0'
            else:
                text = '' if value is None else str(value)
                if ILLEGAL_CHARACTERS_RE.search(text) or len(text) > 32767:
                    raise ExportError('Une valeur ne peut pas être exportée en Excel (texte trop long ou caractères invalides).')
                cell.value = text
                cell.data_type = 's'  # '=...' is literal text, never an Excel formula.
                lines = max(lines, sum(max(1,len(textwrap.wrap(line,width=23))) for line in text.split('\n')))
            cell.font = Font(name='Calibri', size=11, color='17283F')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row_number % 2:
                cell.fill = PatternFill('solid', fgColor='EEF5F8')
        sheet.row_dimensions[row_number].height = min(409, max(30, lines*16))
    for index, key in enumerate(module.columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 12 if key == 'id' else 25
    sheet.freeze_panes = 'A5'
    sheet.auto_filter.ref = f'A4:{get_column_letter(len(module.columns))}{max(4, len(rows)+4)}'
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = '1:4'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_area = sheet.auto_filter.ref
    workbook.save(path)
    workbook.close()


def write_pdf(path, module_key, rows, details):
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import KeepTogether, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    # ReportLab ships Vera: no machine-specific font path or downloaded asset.
    if 'PMVera' not in pdfmetrics.getRegisteredFontNames():
        directory = Path(reportlab.__file__).parent/'fonts'
        pdfmetrics.registerFont(TTFont('PMVera', str(directory/'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('PMVeraBold', str(directory/'VeraBd.ttf')))
    body = ParagraphStyle('PMBody', fontName='PMVera', fontSize=8, leading=11,
                          textColor=colors.HexColor('#17283F'), splitLongWords=True)
    heading = ParagraphStyle('PMHeading', parent=body, fontName='PMVeraBold', textColor=colors.white)
    field_label = ParagraphStyle('PMField', parent=body, fontName='PMVeraBold',
                                 textColor=colors.HexColor('#52647B'))
    title = ParagraphStyle('PMTitle', parent=body, fontName='PMVeraBold', fontSize=18, leading=23)
    note = ParagraphStyle('PMNote', parent=body, fontSize=9, leading=13)
    def paragraph(value, style=body):
        text = '' if value is None else str(value)
        return Paragraph(escape(text).replace('\n', '<br/>'), style)
    document = SimpleDocTemplate(str(path), pagesize=landscape(A4), rightMargin=24, leftMargin=24,
                                 topMargin=26, bottomMargin=32, title=details[0], author='PersonalManager')
    story = [paragraph(details[0], title), Spacer(1, 8),
             paragraph(details[1], note), paragraph(details[2], note), Spacer(1, 16)]
    module = MODULES[module_key]
    wide_or_long = len(module.columns) > 6 or any(len(str(value or '')) > 100 for row in rows for value in row)
    if rows and wide_or_long:
        # A horizontal grid becomes unreadable with many columns/long fields.
        # Use paired fields, splitting between field rows and repeating the ID.
        widths = [document.width*weight for weight in (0.12,0.38,0.12,0.38)]
        for row in rows:
            summary = str(row[1])
            summary = summary if len(summary) <= 80 else summary[:77]+'...'
            name = f'ID {row[0]} · {module.headings[1]} : {summary}'
            data = [[paragraph(name,heading),'','','']]
            fields = list(zip(module.headings[1:],row[1:]))
            for index in range(0,len(fields),2):
                pair = fields[index:index+2]
                cells = []
                for label,value in pair:
                    cells.extend((paragraph(label,field_label),paragraph(value)))
                data.append(cells+['']*(4-len(cells)))
            table = Table(data,colWidths=widths,repeatRows=1,splitInRow=0,hAlign='LEFT')
            table.setStyle(TableStyle([
                ('SPAN',(0,0),(-1,0)),
                ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#142033')),
                ('BACKGROUND',(0,1),(-1,-1),colors.HexColor('#EEF5F8')),
                ('VALIGN',(0,0),(-1,-1),'TOP'),
                ('LEFTPADDING',(0,0),(-1,-1),8), ('RIGHTPADDING',(0,0),(-1,-1),8),
                ('TOPPADDING',(0,0),(-1,-1),8), ('BOTTOMPADDING',(0,0),(-1,-1),8),
            ]))
            story.append(KeepTogether([table,Spacer(1,12)]))
    elif rows:
        weights = [0.5 if key == 'id' else 1.35 if key in ('reason', 'reason_event', 'email', 'meet_with', 'fullname')
                   else 1 for key in module.columns]
        widths = [document.width*weight/sum(weights) for weight in weights]
        data = [[paragraph(value, heading) for value in module.headings]]
        data.extend([paragraph(value) for value in row] for row in rows)
        table = LongTable(data, colWidths=widths, repeatRows=1, splitInRow=0, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#142033')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EEF5F8')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('LINEBELOW', (0, 0), (-1, 0), 1, colors.HexColor('#0F8C9F')),
        ]))
        story.append(table)
    else:
        story.append(paragraph('Aucun enregistrement ne correspond aux critères.', note))
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('PMVera', 8)
        canvas.setFillColor(colors.HexColor('#52647B'))
        canvas.drawString(24, 16, 'PersonalManager · Rapport de données')
        canvas.drawRightString(doc.pagesize[0]-24, 16, f'Page {doc.page}')
        canvas.restoreState()
    document.build(story, onFirstPage=footer, onLaterPages=footer)


def export_report(module_key, destination, actor_id, search='', status=None):
    target = Path(destination)
    suffix = target.suffix.lower()
    if suffix not in ('.xlsx', '.pdf'):
        raise ExportError('Choisissez un fichier .xlsx ou .pdf.')
    rows = export_records(module_key, actor_id, search, status)
    details = report_details(module_key, search, status, len(rows))
    # Build next to the target, then replace only after a successful complete write.
    temporary = None
    try:
        descriptor, temporary = tempfile.mkstemp(prefix='.personalmanager-', suffix=suffix, dir=target.parent)
        os.close(descriptor)
        writer = write_excel if suffix == '.xlsx' else write_pdf
        writer(temporary, module_key, rows, details)
        os.replace(temporary, target)
    except OSError as error:
        raise ExportError('Export impossible. Vérifiez le dossier et fermez le fichier s’il est ouvert.') from error
    finally:
        if temporary is not None:
            Path(temporary).unlink(missing_ok=True)
    return len(rows)
