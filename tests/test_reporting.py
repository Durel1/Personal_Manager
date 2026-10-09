import importlib.util
import unittest
from pathlib import Path
from unittest.mock import patch

from backend.dashboard import export_records
from backend.management import save_record
from backend.permissions import PermissionDenied
from backend.reporting import ExportError, export_report
from test_management import sample
import test_permissions

AVAILABLE = all(importlib.util.find_spec(name) for name in ('openpyxl','reportlab','pypdf'))


@unittest.skipUnless(AVAILABLE, 'Install requirements-dev.txt to verify Excel and PDF exports')
class ReportingTests(unittest.TestCase):
    setUp = test_permissions.PermissionTests.setUp

    def test_excel_exports_all_filtered_pages_and_preserves_text(self):
        from openpyxl import load_workbook
        for i in range(31):
            save_record('clients', dict(sample('clients'), fullname=f'Export {i:02d}'), actor_id=1)
        save_record('clients',dict(sample('clients'),fullname='Other'),actor_id=1)
        target = Path(self.directory.name)/'clients.xlsx'
        self.assertEqual(export_report('clients',target,2,'Export'), 31)
        book = load_workbook(target)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual(sheet.max_row, 35)
        self.assertEqual(sheet['A4'].value, 'ID')
        self.assertEqual(sheet['D5'].value, '00123')
        self.assertEqual(sheet['D5'].data_type, 's')
        self.assertEqual(sheet.freeze_panes, 'A5')
        self.assertEqual(sheet.auto_filter.ref, 'A4:H35')

    def test_excel_never_executes_formulas_and_preserves_large_integer(self):
        from openpyxl import load_workbook
        formula = '=HYPERLINK("https://example.invalid","click")'
        save_record('finances',dict(sample('finances'),reason=formula,amount='9223372036854775807'),actor_id=1)
        target = Path(self.directory.name)/'safe.xlsx'
        export_report('finances',target,1)
        book = load_workbook(target, data_only=False)
        self.addCleanup(book.close)
        self.assertEqual(book.active['B5'].value, formula)
        self.assertEqual(book.active['B5'].data_type, 's')
        self.assertEqual(book.active['C5'].value, '9223372036854775807')

    def test_pdf_repeats_headers_wraps_and_preserves_accents_and_literal_markup(self):
        from pypdf import PdfReader
        for i in range(65):
            save_record('events',dict(sample('events'), meet_with=f'Élise Noël {i}',
                        reason_event='Texte <b>littéral</b> & réunion '+'longue '*45),actor_id=1)
        target = Path(self.directory.name)/'events.pdf'
        self.assertEqual(export_report('events',target,2), 65)
        reader = PdfReader(target)
        self.assertGreater(len(reader.pages), 1)
        all_text = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertIn('Élise Noël 64', all_text)
        self.assertIn('Élise Noël 0', all_text)
        self.assertIn('<b>littéral</b>', all_text)
        for page in reader.pages:
            self.assertIn('Contact',page.extract_text())
            self.assertIn('Page ',page.extract_text())

    def test_empty_reports_and_combined_filters(self):
        from pypdf import PdfReader
        target = Path(self.directory.name)/'empty.pdf'
        self.assertEqual(export_report('clients',target,2,'absent'), 0)
        self.assertIn('Aucun enregistrement', PdfReader(target).pages[0].extract_text())
        for status in ('Payée','Non Payée'):
            save_record('finances',dict(sample('finances'),status=status),actor_id=1)
        target = Path(self.directory.name)/'paid.xlsx'
        self.assertEqual(export_report('finances',target,1,'Facture','Payée'), 1)

    def test_permission_and_extension_errors_do_not_create_files(self):
        target = Path(self.directory.name)/'private.xlsx'
        with self.assertRaises(PermissionDenied):
            export_report('finances',target,2)
        self.assertFalse(target.exists())
        with self.assertRaises(ExportError):
            export_report('clients',target.with_suffix('.csv'),1)

    def test_failure_preserves_existing_file_and_removes_temporary_file(self):
        target = Path(self.directory.name)/'report.xlsx'
        target.write_bytes(b'previous report')
        def broken_writer(path,*args):
            Path(path).write_bytes(b'partial')
            raise OSError('simulated locked disk')
        with patch('backend.reporting.write_excel',broken_writer), self.assertRaises(ExportError):
            export_report('clients',target,1)
        self.assertEqual(target.read_bytes(),b'previous report')
        self.assertEqual(list(target.parent.glob('.personalmanager-*')),[])

    def test_limit_is_explicit_and_never_silently_truncates(self):
        for _ in range(3):
            save_record('clients',sample('clients'),actor_id=1)
        with self.assertRaises(ValueError):
            export_records('clients',1,limit=2)
