"""Real Windows widget smoke tests; layout still needs human visual review."""
import importlib.util
import os
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

AVAILABLE = all(importlib.util.find_spec(name) is not None for name in ('bcrypt', 'customtkinter', 'matplotlib'))


@unittest.skipUnless(AVAILABLE, 'bcrypt/customtkinter/matplotlib unavailable: modern widget checks pending')
class ModernWidgetTests(unittest.TestCase):
    def setUp(self):
        # Unattended tests must never wait for a user to dismiss an error dialog.
        dialogs = patch('ui.application.messagebox.showerror')
        self.error_dialog = dialogs.start()
        self.addCleanup(dialogs.stop)
        self.callback_errors = []
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        env = patch.dict(os.environ, {'PERSONAL_MANAGER_DB': str(Path(directory.name)/'widgets.db')})
        env.start()
        self.addCleanup(env.stop)
        from ui.application import PersonalManager
        self.app = PersonalManager()
        self.addCleanup(self.shutdown)
        original_report = self.app.report_callback_exception

        def capture_callback_error(exception_type, error, traceback):
            self.callback_errors.append((error, traceback))
            original_report(exception_type, error, traceback)

        self.app.report_callback_exception = capture_callback_error
        self.wait_for(lambda: self.app.generation >= 2)

    def shutdown(self):
        # Poll completion before releasing the temporary directory on Windows.
        self.app.executor.shutdown(wait=True, cancel_futures=True)
        self.app.close()

    def wait_for(self, condition):
        deadline = time.monotonic()+5
        while time.monotonic() < deadline:
            self.app.update()
            if self.callback_errors:
                error, traceback = self.callback_errors.pop(0)
                raise error.with_traceback(traceback)
            if condition():
                return
            time.sleep(0.01)
        self.fail('Modern UI operation did not finish within five seconds')

    def drain_tasks(self):
        self.wait_for(lambda: not self.app.polls and self.app.search_timer is None)

    def test_login_and_registration_render_in_one_root(self):
        root_id = str(self.app)
        self.app.show_register()
        self.app.update()
        self.app.show_login()
        self.app.update()
        self.assertEqual(str(self.app), root_id)
        self.assertIsNone(self.app.user)

    def test_sidebar_lists_theme_and_logout(self):
        from backend.auth import register_user
        identifier = register_user('UITest', 'Secret123!', 'test@example.com', '00123', 'Homme')
        self.app.user = {'id': identifier, 'fullname': 'UITest'}
        self.app.show_shell()
        self.drain_tasks()
        for key in ('employees', 'clients', 'events', 'finances', 'profile'):
            self.app.navigate(key)
            self.drain_tasks()
        self.app.change_theme('Clair')
        self.app.update()
        self.app.change_theme('Sombre')
        self.app.update()
        self.app.show_login()
        self.assertIsNone(self.app.user)

    def test_navigation_ignores_obsolete_worker_result(self):
        completed = []
        self.app.run_task(lambda: 42, completed.append)
        self.app.show_register()
        self.drain_tasks()
        self.assertEqual(completed, [])

    def test_close_cancels_library_and_child_widget_callbacks(self):
        completed = []
        errors = []
        self.app.screen.after(1000,lambda: completed.append('child'))
        self.app.after_idle(lambda: completed.append('idle'))
        self.app.tk.createcommand('pm_test_bgerror',errors.append)
        self.app.tk.eval('proc bgerror {message} {pm_test_bgerror $message}')
        try:
            self.app.close()
            self.assertEqual(self.app.tk.splitlist(self.app.tk.call('after','info')), ())
            self.app.tk.call('update')
            self.assertEqual(completed, [])
            self.assertEqual(errors, [])
            self.app.close()  # Repeated close is harmless.
        finally:
            self.app.tk.deletecommand('pm_test_bgerror')

    def test_failed_success_callback_displays_error_instead_of_crashing(self):
        def broken(_):
            raise RuntimeError('simulated callback failure')
        with patch('ui.application.messagebox.showerror') as dialog, patch('ui.application.record_error'):
            self.app.run_task(lambda: 42,broken)
            with self.assertRaisesRegex(RuntimeError, 'simulated callback failure'):
                self.drain_tasks()
            dialog.assert_called_once()

    def sign_in_test_user(self):
        from backend.auth import register_user
        identifier = register_user('FormTest', 'Secret123!', 'test@example.com', '00123', 'Homme')
        self.app.user = {'id': identifier, 'fullname': 'FormTest'}
        self.app.show_shell()
        self.drain_tasks()

    def fill_finance_form(self, amount):
        values = dict(reason='UI invoice', amount=amount, date='2026-10-09',
                      status='Non Payée', type='Encaissement')
        for key, value in values.items():
            widget = self.app.form_fields[key]
            if key in ('status', 'type'):
                widget.set(value)
            else:
                widget.delete(0, 'end')
                widget.insert(0, value)

    def test_finance_form_creates_then_updates_same_record(self):
        from backend.management import get_record
        from backend.dashboard import dashboard_counts
        self.sign_in_test_user()
        self.app.show_form('finances')
        self.fill_finance_form('1000')
        self.app.form_save_button.invoke()
        self.drain_tasks()
        self.assertEqual(dashboard_counts(actor_id=self.app.user['id'])['finances'], 1)
        identifier = int(self.app.record_tree.get_children()[0])
        self.app.show_form('finances', identifier)
        self.drain_tasks()
        self.fill_finance_form('2500')
        self.app.form_save_button.invoke()
        self.drain_tasks()
        self.assertEqual(dashboard_counts(actor_id=self.app.user['id'])['finances'], 1)
        self.assertEqual(get_record('finances', identifier, actor_id=self.app.user['id'])['amount'], 2500)

    def test_invalid_form_shows_message_without_writing(self):
        from backend.dashboard import dashboard_counts
        self.sign_in_test_user()
        self.app.show_form('finances')
        self.fill_finance_form('not-a-number')
        self.app.form_save_button.invoke()
        self.drain_tasks()
        self.assertIn('entier positif', self.app.form_status.cget('text'))
        self.assertEqual(dashboard_counts(actor_id=self.app.user['id'])['finances'], 0)

    def test_dashboard_charts_follow_theme_and_are_released_on_navigation(self):
        self.sign_in_test_user()
        self.assertEqual(len(self.app.chart_canvases), 2)
        old_chart = self.app.chart_canvases[-1].get_tk_widget()
        old_chart.focus_set()
        self.app.update()
        self.app.change_theme('Clair')
        self.drain_tasks()
        self.assertFalse(old_chart.winfo_exists())
        self.assertEqual(len(self.app.chart_canvases), 2)
        self.app.navigate('clients')
        self.drain_tasks()
        self.assertEqual(self.app.chart_canvases, [])
        self.app.navigate('home')
        self.drain_tasks()
        self.assertEqual(len(self.app.chart_canvases), 2)

    def test_live_search_keeps_criteria_after_form_navigation(self):
        from backend.management import save_record
        self.sign_in_test_user()
        for name in ('Alice', 'Bob'):
            save_record('employees',dict(fullname=name,email='test@example.com',phone='00123',gender='Femme'),
                        actor_id=self.app.user['id'])
        self.app.navigate('employees')
        self.drain_tasks()
        self.app.search_entry.insert(0, 'alice')
        self.drain_tasks()
        rows = self.app.record_tree.get_children()
        self.assertEqual(len(rows), 1)
        identifier = int(rows[0])
        self.app.show_form('employees',identifier)
        self.drain_tasks()
        self.app.navigate('employees')
        self.drain_tasks()
        self.assertEqual(self.app.search_entry.get(), 'alice')
        self.assertEqual(len(self.app.record_tree.get_children()), 1)

    def test_search_status_filter_and_reset(self):
        from backend.management import save_record
        self.sign_in_test_user()
        for state in ('Payée','Non Payée'):
            save_record('finances',dict(reason='Loyer',amount='1000',date='2026-10-09',status=state,type='Décaissement'),
                        actor_id=self.app.user['id'])
        self.app.navigate('finances')
        self.drain_tasks()
        self.app.status_filter.set('Payée')
        self.app.refresh_search()
        self.drain_tasks()
        self.assertEqual(len(self.app.record_tree.get_children()), 1)
        self.app.status_filter.set('Tous')
        self.app.refresh_search()
        self.drain_tasks()
        self.assertEqual(len(self.app.record_tree.get_children()), 2)

    def test_employee_shell_has_only_authorized_views_and_readonly_actions(self):
        from backend.auth import register_user
        self.sign_in_test_user()
        identifier = register_user('Viewer','Secret123!','viewer@example.com','002','Femme')
        self.app.user = {'id':identifier,'fullname':'Viewer'}
        self.app.show_shell()
        self.drain_tasks()
        self.assertNotIn('finances',self.app.navigation)
        self.assertNotIn('employees',self.app.navigation)
        self.assertNotIn('users',self.app.navigation)
        self.assertEqual(len(self.app.chart_canvases),1)
        with patch('ui.application.messagebox.showerror') as error:
            self.app.navigate('finances')
            error.assert_called_once()
        self.assertEqual(self.app.current_view,'home')
        self.app.navigate('clients')
        self.drain_tasks()
        self.assertEqual(self.app.record_add_button.cget('state'),'disabled')
        self.assertEqual(self.app.record_delete_button.cget('state'),'disabled')
        self.assertEqual(self.app.excel_export_button.cget('state'),'normal')

    def test_export_button_writes_filtered_rows_and_cancel_writes_nothing(self):
        from backend.management import save_record
        from openpyxl import load_workbook
        self.sign_in_test_user()
        for name in ('Alice','Bob'):
            save_record('employees',dict(fullname=name,email='test@example.com',phone='00123',gender='Femme'),
                        actor_id=self.app.user['id'])
        self.app.navigate('employees')
        self.drain_tasks()
        self.app.search_entry.insert(0,'alice')
        self.drain_tasks()
        target = Path(os.environ['PERSONAL_MANAGER_DB']).with_suffix('.xlsx')
        with patch('ui.application.filedialog.asksaveasfilename',return_value=str(target)):
            self.app.excel_export_button.invoke()
            self.drain_tasks()
        book = load_workbook(target)
        try:
            self.assertEqual(book.active.max_row,5)
            self.assertEqual(book.active['B5'].value,'Alice')
        finally:
            book.close()
        with patch('ui.application.filedialog.asksaveasfilename',return_value=''):
            self.app.pdf_export_button.invoke()
        self.assertFalse(target.with_suffix('.pdf').exists())

    def test_today_and_overdue_rows_receive_alert_tags(self):
        from datetime import date,timedelta
        from backend.management import save_record
        self.sign_in_test_user()
        today = date.today()
        save_record('events',dict(meet_with='Client',gender='Homme',phone='00123',place='Bureau',
                    event_status='Non Effectué',reason_event='Réunion',eventdate=today.isoformat(),hour_event='09:30'),
                    actor_id=self.app.user['id'])
        save_record('finances',dict(reason='Facture',amount='1000',date=(today-timedelta(days=2)).isoformat(),
                    status='Non Payée',type='Encaissement',due_date=(today-timedelta(days=1)).isoformat()),
                    actor_id=self.app.user['id'])
        self.app.navigate('home')
        self.drain_tasks()
        self.assertIn('Factures en retard : 1',self.app.alert_label.cget('text'))
        for key,tag in (('events','today'),('finances','overdue')):
            self.app.navigate(key)
            self.drain_tasks()
            row = self.app.record_tree.get_children()[0]
            self.assertIn(tag,self.app.record_tree.item(row,'tags'))

    def test_admin_assigns_role_in_permissions_view(self):
        from backend.auth import register_user
        from backend.permissions import current_user
        self.sign_in_test_user()
        identifier = register_user('NewUser','Secret123!','new@example.com','002','Homme')
        self.app.navigate('users')
        self.drain_tasks()
        self.assertEqual(len(self.app.user_tree.get_children()),2)
        self.app.user_tree.selection_set(str(identifier))
        self.app.role_choice.set('Administrateur')
        with patch('ui.application.messagebox.askyesno',return_value=True):
            self.app.role_save_button.invoke()
            self.drain_tasks()
        self.assertEqual(current_user(identifier)['role'],'admin')
